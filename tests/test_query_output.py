import json

from openpyxl import load_workbook

from app.query_output import jsonl_to_xlsx


def test_jsonl_to_xlsx_flattens_objects_and_preserves_lists_as_json(tmp_path):
    jsonl = tmp_path / "result.jsonl"
    xlsx = tmp_path / "result.xlsx"
    jsonl.write_text(
        json.dumps(
            {
                "id": "gid://shopify/MediaImage/1",
                "image": {"url": "https://example.com/image.jpg"},
                "tags": ["A", "B"],
                "__parentId": "gid://shopify/Product/2",
            }
        )
        + "\n",
        encoding="utf-8",
    )
    jsonl_to_xlsx(jsonl, xlsx)
    workbook = load_workbook(xlsx, read_only=True)
    sheet = workbook["Query Result"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("id", "image.url", "tags", "__parentId")
    assert rows[1] == (
        "gid://shopify/MediaImage/1",
        "https://example.com/image.jpg",
        '["A","B"]',
        "gid://shopify/Product/2",
    )
