from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.utils import iter_jsonl

EXCEL_CELL_LIMIT = 32767


def _flatten(value: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    row: dict[str, Any] = {}
    for key, item in value.items():
        column = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, dict):
            row.update(_flatten(item, column))
        elif isinstance(item, list):
            row[column] = json.dumps(item, ensure_ascii=False, separators=(",", ":"))
        else:
            row[column] = item
    return row


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    return text[:EXCEL_CELL_LIMIT]


def jsonl_to_xlsx(jsonl_path: Path, xlsx_path: Path) -> Path:
    """Create a streaming XLSX convenience view while leaving the JSONL authoritative."""
    columns: list[str] = []
    seen: set[str] = set()
    row_count = 0
    for obj in iter_jsonl(jsonl_path):
        row_count += 1
        for column in _flatten(obj):
            if column not in seen:
                seen.add(column)
                columns.append(column)

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Query Result")
    sheet.freeze_panes = "A2"

    if columns:
        sheet.append(columns)
        for obj in iter_jsonl(jsonl_path):
            flat = _flatten(obj)
            sheet.append([_excel_value(flat.get(column)) for column in columns])
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row_count + 1}"

    workbook.save(xlsx_path)
    return xlsx_path
